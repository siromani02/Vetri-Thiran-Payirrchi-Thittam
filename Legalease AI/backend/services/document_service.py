from io import BytesIO
import re
from pathlib import Path
from xml.sax.saxutils import escape
from PIL import Image as PILImage
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
ROOT_DIR = Path(__file__).resolve().parents[2]
LOGO_PATH = ROOT_DIR / "Image" / "Logo.png"
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (
    Image as RLImage,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


# ============================================================
# TEXT HELPERS
# ============================================================

def clean_text(text):
    text = text or ""

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Remove markdown headings
    text = re.sub(
        r"^\s*#{1,6}\s*",
        "",
        text,
        flags=re.MULTILINE,
    )

    # Remove markdown bold markers
    text = text.replace("**", "")
    text = text.replace("__", "")

    return text.strip()


def split_lines(text):
    return [
        line.strip()
        for line in clean_text(text).split("\n")
        if line.strip()
    ]


def is_heading(line):
    if not line:
        return False

    # Numbered headings such as:
    # 1. Services
    # 2. Payment
    if re.match(r"^\d+[\.)]\s+", line):
        return True

    # Fully uppercase headings
    letters = re.sub(r"[^A-Za-z]", "", line)

    if (
        letters
        and letters.upper() == letters
        and len(letters) >= 4
    ):
        return True

    known_headings = {
        "parties",
        "effective date",
        "jurisdiction",
        "terms and conditions",
        "services",
        "term and termination",
        "payment",
        "intellectual property rights",
        "confidentiality",
        "independent contractor status",
        "governing law",
        "entire agreement",
        "severability",
        "termination",
        "signatures",
        "signature",
        "witnesseth",
        "whereas",
        "now therefore",
        "in witness whereof",
    }

    return line.lower().strip() in known_headings


# ============================================================
# FOOTER FOR PDF
# ============================================================

def pdf_footer(canvas, doc):
    canvas.saveState()

    width, height = A4

    canvas.setStrokeColor(
        colors.HexColor("#AAAAAA")
    )

    canvas.setLineWidth(0.5)

    canvas.line(
        20 * mm,
        15 * mm,
        width - 20 * mm,
        15 * mm,
    )

    canvas.setFont(
        "Helvetica",
        8,
    )

    canvas.setFillColor(
        colors.HexColor("#666666")
    )

    canvas.drawString(
        20 * mm,
        10 * mm,
        "LegalEase - AI-Powered Legal Document Generator",
    )

    canvas.drawRightString(
        width - 20 * mm,
        10 * mm,
        f"Page {doc.page}",
    )

    canvas.restoreState()


# ============================================================
# TXT
# ============================================================

def create_txt(text):
    content = clean_text(text)

    final_text = f"""
============================================================
                         LegalEase
             AI-Powered Legal Document Generator
============================================================

{content}

------------------------------------------------------------
IMPORTANT LEGAL NOTICE

This document is an AI-generated draft for informational
purposes only. It is not legal advice.

Consult a qualified legal professional before using or
signing this document.
------------------------------------------------------------
"""

    return final_text.encode("utf-8")


# ============================================================
# PDF
# ============================================================

def create_pdf(text):
    buffer = BytesIO()

    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=38 * mm,
        bottomMargin=22 * mm,
        title="LegalEase Legal Document",
        author="LegalEase",
    )

    styles = getSampleStyleSheet()

    # --------------------------------------------------------
    # Brand
    # --------------------------------------------------------

    brand_style = ParagraphStyle(
        "Brand",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        alignment=TA_CENTER,
        textColor=colors.black,
        spaceAfter=2 * mm,
    )

    subtitle_style = ParagraphStyle(
        "Subtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#555555"),
        spaceAfter=5 * mm,
    )

    # --------------------------------------------------------
    # Main title
    # --------------------------------------------------------

    title_style = ParagraphStyle(
        "Title",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=17,
        leading=21,
        alignment=TA_CENTER,
        textColor=colors.black,
        spaceBefore=5 * mm,
        spaceAfter=8 * mm,
    )

    # --------------------------------------------------------
    # Section heading
    # --------------------------------------------------------

    heading_style = ParagraphStyle(
        "Heading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=15,
        alignment=TA_LEFT,
        textColor=colors.black,
        spaceBefore=5 * mm,
        spaceAfter=2 * mm,
        keepWithNext=True,
    )

    # --------------------------------------------------------
    # Body
    # --------------------------------------------------------

    body_style = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        alignment=TA_LEFT,
        textColor=colors.black,
        spaceAfter=3.5 * mm,
    )

    story = []

    # ========================================================
    # LEGAL EASE BRANDING
    # ========================================================
    # Horizontal line
    line = Table(
        [[""]],
        colWidths=[170 * mm],
        rowHeights=[0.5 * mm],
    )

    line.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, -1),
                    colors.black,
                )
            ]
        )
    )

    story.append(line)

    story.append(
        Spacer(
            1,
            4 * mm,
        )
    )

    # ========================================================
    # FIND DOCUMENT TITLE
    # ========================================================

    lines = split_lines(text)

    title_text = "LEGAL DOCUMENT"

    for line_text in lines[:10]:

        lower = line_text.lower()

        if any(
            word in lower
            for word in [
                "contract",
                "agreement",
                "nda",
                "non-disclosure",
                "lease",
                "employment",
                "freelance",
            ]
        ):
            title_text = line_text.upper()
            break

    story.append(
        Paragraph(
            escape(title_text),
            title_style,
        )
    )

    # ========================================================
    # MAIN CONTENT
    # ========================================================

    title_skipped = False

    for line_text in lines:

        lower = line_text.lower().strip()

        # Skip duplicate title
        if not title_skipped:
            if lower == title_text.lower():
                title_skipped = True
                continue

            title_skipped = True

        # Skip branding text if Gemini included it
        if lower in {
            "legalease",
            "ai-powered legal document generator",
        }:
            continue

        # Skip old warning if Gemini includes it
        if lower.startswith("important legal notice"):
            continue

        # Make headings bold
        if is_heading(line_text):

            story.append(
                Paragraph(
                    escape(line_text),
                    heading_style,
                )
            )

        else:

            story.append(
                Paragraph(
                    escape(line_text),
                    body_style,
                )
            )

    # ========================================================
    # SIGNATURES
    # ========================================================

    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    story.append(
        Paragraph(
            "SIGNATURES",
            heading_style,
        )
    )

    signature_table = Table(
        [
            [
                Paragraph(
                    "<b>Service Provider</b>",
                    body_style,
                ),
                Paragraph(
                    "<b>Client</b>",
                    body_style,
                ),
            ],
            [
                Paragraph(
                    "Signature: __________________________<br/>"
                    "Name: _______________________________<br/>"
                    "Date: ________________________________",
                    body_style,
                ),
                Paragraph(
                    "Signature: __________________________<br/>"
                    "Name: _______________________________<br/>"
                    "Date: ________________________________",
                    body_style,
                ),
            ],
        ],
        colWidths=[
            82 * mm,
            82 * mm,
        ],
    )

    signature_table.setStyle(
        TableStyle(
            [
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    0,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    5 * mm,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    2,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    3,
                ),
            ]
        )
    )

    story.append(signature_table)

    # ========================================================
    # BUILD PDF
    # ========================================================S
    pdf.build(
        story,
        onFirstPage=pdf_header_footer,
        onLaterPages=pdf_header_footer,
    )

    buffer.seek(0)

    return buffer.getvalue()
def pdf_header(canvas, doc):
    canvas.saveState()

    if LOGO_PATH.exists():
        logo_width = 48 * mm
        logo_height = 18 * mm

        x = (A4[0] - logo_width) / 2
        y = A4[1] - 30 * mm

        canvas.drawImage(
            str(LOGO_PATH),
            x,
            y,
            width=logo_width,
            height=logo_height,
            preserveAspectRatio=True,
            mask="auto",
        )

    canvas.restoreState()

def pdf_header_footer(canvas, doc):
    pdf_header(canvas, doc)
    pdf_footer(canvas, doc)


# ============================================================
# DOCX
# ============================================================

def create_docx(
    text,
    document_type="Legal Document",
):

    document = Document()

    section = document.sections[0]

    # --------------------------------------------------------
    # Page margins
    # --------------------------------------------------------

    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # --------------------------------------------------------
    # Default font
    # --------------------------------------------------------

    normal = document.styles["Normal"]

    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)

    # ========================================================
    # HEADER
    # ========================================================

    header = section.header

    paragraph = header.paragraphs[0]

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "⚖ LegalEase"
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)

    paragraph = header.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "AI-Powered Legal Document Generator"
    )

    run.font.name = "Times New Roman"
    run.font.size = Pt(9)

    # ========================================================
    # TITLE
    # ========================================================

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        document_type.upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(17)

    paragraph.paragraph_format.space_after = Pt(10)

    # ========================================================
    # CONTENT
    # ========================================================

    lines = split_lines(text)

    title_lower = document_type.lower()

    for line_text in lines:

        lower = line_text.lower().strip()

        # Skip duplicate title
        if lower == title_lower:
            continue

        # Skip branding if present
        if lower in {
            "legalease",
            "ai-powered legal document generator",
        }:
            continue

        # Skip old warning
        if lower.startswith("important legal notice"):
            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15

        # Numbered sections / headings
        if is_heading(line_text):

            paragraph.paragraph_format.space_before = Pt(8)

            run = paragraph.add_run(
                line_text
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

        else:

            run = paragraph.add_run(
                line_text
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(10)

    # ========================================================
    # SIGNATURES
    # ========================================================

    document.add_page_break()

    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        "The parties acknowledge that they have read, "
        "understood, and agreed to the terms of this Agreement "
        "and have executed it as of the Effective Date."
    )

    run.font.name = "Times New Roman"
    run.font.size = Pt(10)

    document.add_paragraph()

    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        "SIGNATURES"
    )


    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    paragraph.paragraph_format.space_after = Pt(8)

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.autofit = True

    # Headers
    table.cell(
        0,
        0,
    ).text = "Service Provider"

    table.cell(
        0,
        1,
    ).text = "Client"

    # Signature fields
    table.cell(
        1,
        0,
    ).text = (
        "\n\n"
        "Signature: ______________________\n"
        "Name: ___________________________\n"
        "Date: ____________________________"
    )

    table.cell(
        1,
        1,
    ).text = (
        "\n\n"
        "Signature: ______________________\n"
        "Name: ___________________________\n"
        "Date: ____________________________"
    )

    # Format table text
    for row in table.rows:

        for cell in row.cells:

            for paragraph in cell.paragraphs:

                for run in paragraph.runs:

                    run.font.name = "Times New Roma"
                    run.font.size = Pt(10)

    # ========================================================
    # FOOTER
    # ========================================================

    footer = section.footer

    paragraph = footer.paragraphs[0]

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "LegalEase | AI-Powered Legal Document Generator | All Rights Reserved"
    )

    run.font.name = "Times New Roman"
    run.font.size = Pt(8)

    # ========================================================
    # SAVE DOCX TO MEMORY
    # ========================================================

    buffer = BytesIO()

    document.save(buffer)

    buffer.seek(0)

    return buffer.getvalue()